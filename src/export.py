"""Exports a ResearchFlow report (markdown from report_builder) to md/docx/pdf.

The docx and pdf writers understand the subset of markdown the report uses: #-#### headings,
bullets, *italic* lines, > notes, **bold** spans, pipe tables, ![images](path) relative to the
output directory, and the critical-thinking (and Phase 2-4) markers, which are highlighted so
they stand out.
"""
from __future__ import annotations

import logging
import re
from pathlib import Path
from typing import List

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Inches, Pt
from fpdf import FPDF
from fpdf.enums import XPos, YPos

from .citation_formatter import CitationFormatter
from .critical_gaps import HIGHLIGHT_PATTERN
from .paper_discovery import Paper

logger = logging.getLogger(__name__)

IMAGE_LINE = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$")
BOLD_SPAN = re.compile(r"(\*\*[^*]+\*\*)")
# Characters outside latin-1 that the core PDF fonts can't draw, with readable substitutes.
PDF_REPLACEMENTS = str.maketrans({
    "“": '"', "”": '"', "‘": "'", "’": "'", "–": "-", "—": "-", "…": "...", "•": "-",
    # statistical symbols in the Phase 3 results
    "α": "alpha", "β": "beta", "κ": "kappa", "η": "eta", "χ": "chi", "ρ": "rho", "≤": "<=", "≥": ">=",
})
TABLE_COL_WIDTHS = (34, 22, 10, 50, 46, 28)  # mm, fits the 190 mm text width of an A4 page
TABLE_FIRST_HEADER = "Paper Title"  # the widths are for the Results & Analysis table; other tables share the width evenly


# ---- executive summary --------------------------------------------------------------------------

SUMMARY_MAX_WORDS = 550  # about one page
SUMMARY_PARTS = (  # (label, conclusion heading it is drawn from, sentences taken)
    ("What the field knows", "1.1", 3),
    ("How the proposed study builds on it", "1.2", 2),
    ("What the analysis found", "1.3", 2),
    ("Contribution", "1.4", 2),
    ("Limitations", "1.5", 2),
)


def _subsections(markdown: str) -> dict:
    """{"1.1": [prose lines], ...} for the conclusion's numbered "### n.n" subsections."""
    sections: dict = {}
    current = None
    for line in markdown.splitlines():
        heading = re.match(r"^###\s+(\d+\.\d+)\s", line)
        if heading:
            current = heading.group(1)
            sections[current] = []
        elif line.startswith("## "):
            current = None
        elif current and line.strip():
            sections[current].append(line.strip())
    return sections


def _prose_lines(lines: List[str]) -> List[str]:
    """Body text only: no marker prompts, guidance notes, tables or blockquotes."""
    return [
        line for line in lines
        if not line.startswith(("|", ">", "**[", "*[")) and not (line.startswith("*") and line.endswith("*") and not line.startswith("**"))
    ]


def build_executive_summary(conclusion: str, topic: str, degree: str = "") -> str:
    """A one-page executive summary drawn from a Phase 4 conclusion (no new claims are added).

    Takes the opening sentences of each Conclusion subsection, the practical guidance, the
    highest-priority future directions and the places the user must still fill, and closes with the
    draft's readiness for the chosen degree (thesis_standards).
    """
    from .critical_gaps import CONCLUSION_MARKER_LINE, USER_MUST_ADD, guidance
    from .llm_synthesis import split_sentences
    from .thesis_standards import PHASE_READINESS, degree_standard

    sections = _subsections(conclusion)
    standard = degree_standard(degree)
    lines = [
        f"# Executive Summary: {topic}",
        "",
        "*A one-page summary of the conclusion, implications and future directions, for supervisors, committees and busy readers*",
        "",
    ]
    for label, key, count in SUMMARY_PARTS:
        # Lists of several citations are dropped: a one-page summary names findings, not every source.
        prose = re.sub(r"\s*\((?:[^()]*\b(?:\d{4}|n\.d\.)\s*;){1,}[^()]*\)", "", " ".join(_prose_lines(sections.get(key, []))))
        sentences = split_sentences(prose)
        if key == "1.5":
            # The review's own limitations matter most to a reader of the summary.
            own = [s for s in sentences if s.startswith(("The review itself", "Papers were selected", "Only ", "None of the abstracts"))]
            sentences = own or sentences
        sentences = sentences[:count]
        if sentences:
            lines += [f"**{label}.** {' '.join(sentences)}", ""]

    practice = [re.sub(r"^\d+\.\s*", "", l) for l in sections.get("2.5", []) if re.match(r"^\d+\.\s", l)]
    if not practice:
        practice = [l[2:] for l in sections.get("2.1", []) if l.startswith("- ")]
    if practice:
        lines += ["**For practitioners.**", ""] + [f"- {split_sentences(p)[0] if split_sentences(p) else p}" for p in practice[:3]] + [""]

    rows = [l for l in conclusion.splitlines() if l.startswith("| ") and not l.startswith(("| Direction", "|---"))]
    directions = []
    for row in rows:
        cells = [c.strip() for c in row.strip("|").split("|")]
        if len(cells) == 5:
            title, kind, _, timeline, feasibility = cells
            directions.append((title, kind, timeline, feasibility))
    if directions:
        ranked = sorted(directions, key=lambda d: {"High": 0, "Medium": 1, "Low": 2}.get(d[3], 3))[:3]
        lines += ["**Priority directions.**", ""]
        lines += [f"- {title} ({kind.lower()}; {timeline[:1].lower() + timeline[1:]}; feasibility {feasibility.lower()})"
                  for title, kind, timeline, feasibility in ranked]
        lines.append("")

    prompts = []
    for line in conclusion.splitlines():
        match = CONCLUSION_MARKER_LINE.match(line)
        if match:
            prompt = line[match.end():].strip()
            prompts.append(f"{match.group(1)} {split_sentences(prompt)[0] if split_sentences(prompt) else prompt}")
    if prompts:
        lines += ["**What only you can add.**", ""] + [f"- {p}" for p in dict.fromkeys(prompts)] + [""]

    lines += [
        f"**Readiness for {standard.name}.** ResearchFlow targets {standard.readiness}% of an approval-ready draft at "
        f"this level (Phases 1-4: " + ", ".join(f"{phase.split(' (')[0]} {pct}%" for phase, pct in PHASE_READINESS.items())
        + f"). The remaining {100 - standard.readiness}% is your own data, analysis and argument, which examiners at this "
        f"level judge on {', '.join(standard.examiner_focus[:2])} and {standard.examiner_focus[2]}.",
        "",
        guidance(USER_MUST_ADD, "Replace the summary of findings with your own results once your study is complete."),
        "",
    ]
    summary = "\n".join(lines)
    words = len(re.findall(r"\w+", summary))
    if words > SUMMARY_MAX_WORDS:
        logger.info("Executive summary is %d words (target %d); trimming the longest parts", words, SUMMARY_MAX_WORDS)
        summary = _trim_summary(summary)
    return summary


def _trim_summary(summary: str) -> str:
    """Shortens each labelled paragraph to its first sentence until the summary fits about one page."""
    from .llm_synthesis import split_sentences

    out = []
    for line in summary.splitlines():
        match = re.match(r"^(\*\*[^*]+\.\*\*) (.+)$", line)
        if match and len(re.findall(r"\w+", line)) > 45 and not line.startswith("**Readiness"):
            first = split_sentences(match.group(2))[:1]
            line = f"{match.group(1)} {first[0]}" if first else line
        out.append(line)
    return "\n".join(out)


def _table_rows(lines: List[str]) -> List[List[str]]:
    """Cells of a markdown pipe table, without its |---| separator row."""
    rows = []
    for line in lines:
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
            continue
        rows.append(cells)
    return rows


def _blocks(content: str):
    """Yields ("table", [lines]) for runs of table lines and ("line", line) for everything else."""
    table: List[str] = []
    for line in content.splitlines():
        if line.lstrip().startswith("|"):
            table.append(line)
            continue
        if table:
            yield "table", table
            table = []
        yield "line", line
    if table:
        yield "table", table


class Exporter:
    def __init__(self, output_dir: Path):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def export_all(self, content: str, slug: str, formats: List[str]) -> List[Path]:
        """Writes the report in each requested format, markdown first.

        A failure in one format is logged and skipped so it can't prevent the others; markdown is
        always written (even if not requested) when another format fails, so the run's output is
        never lost.
        """
        exporters = [
            ("md", self.export_markdown),
            ("docx", self.export_docx),
            ("pdf", self.export_pdf),
        ]
        paths: List[Path] = []
        failed = False
        for fmt, export in exporters:
            if fmt not in formats:
                continue
            try:
                paths.append(export(content, f"{slug}.{fmt}"))
            except Exception as exc:
                failed = True
                logger.error("Failed to export %s: %s", fmt, exc)
        if failed and "md" not in formats:
            paths.append(self.export_markdown(content, f"{slug}.md"))
            logger.info("Wrote markdown fallback because another format failed")
        return paths

    def export_markdown(self, content: str, filename: str) -> Path:
        path = self.output_dir / filename
        path.write_text(content, encoding="utf-8")
        return path

    def _image_path(self, target: str) -> Path:
        path = Path(target)
        return path if path.is_absolute() else self.output_dir / path

    # ---- docx -----------------------------------------------------------------------------------

    @staticmethod
    def _add_runs(paragraph, text: str, italic: bool = False) -> None:
        """Adds text with **bold** spans; critical-thinking markers are bold and highlighted."""
        for part in BOLD_SPAN.split(text):
            if not part:
                continue
            bold = part.startswith("**") and part.endswith("**")
            part = part[2:-2] if bold else part
            for piece in re.split(f"({HIGHLIGHT_PATTERN.pattern})", part):
                if not piece:
                    continue
                run = paragraph.add_run(piece)
                run.bold = bold or bool(HIGHLIGHT_PATTERN.fullmatch(piece))
                run.italic = italic
                if HIGHLIGHT_PATTERN.fullmatch(piece):
                    run.font.highlight_color = WD_COLOR_INDEX.YELLOW

    def export_docx(self, content: str, filename: str) -> Path:
        document = Document()
        document.styles["Normal"].font.size = Pt(11)

        for kind, block in _blocks(content):
            if kind == "table":
                self._docx_table(document, _table_rows(block))
                continue
            line = block
            stripped = line.strip()
            heading = re.match(r"^(#{1,4}) (.*)$", line)
            image = IMAGE_LINE.match(stripped)
            if heading:
                document.add_heading(heading.group(2).replace("**", ""), level=len(heading.group(1)))
            elif image:
                path = self._image_path(image.group(2))
                if path.exists():
                    document.add_picture(str(path), width=Inches(6.0))
                else:
                    logger.warning("Image %s not found; leaving it out of the docx", path)
            elif line.startswith("- "):
                self._add_runs(document.add_paragraph(style="List Bullet"), line[2:])
            elif stripped.startswith("> "):
                self._add_runs(document.add_paragraph(), stripped[2:], italic=True)
            elif stripped.startswith("*") and stripped.endswith("*") and not stripped.startswith("**"):
                self._add_runs(document.add_paragraph(), stripped.strip("*"), italic=True)
            elif stripped:
                self._add_runs(document.add_paragraph(), line)

        path = self.output_dir / filename
        document.save(str(path))
        return path

    def _docx_table(self, document, rows: List[List[str]]) -> None:
        if not rows:
            return
        columns = max(len(r) for r in rows)
        table = document.add_table(rows=len(rows), cols=columns)
        table.style = "Table Grid"
        for r, cells in enumerate(rows):
            for c, text in enumerate(cells):
                paragraph = table.cell(r, c).paragraphs[0]
                run = paragraph.add_run(text.replace("**", ""))
                run.font.size = Pt(8)
                run.bold = r == 0
        document.add_paragraph()

    # ---- pdf ------------------------------------------------------------------------------------

    def export_pdf(self, content: str, filename: str) -> Path:
        pdf = FPDF()
        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.add_page()

        cell_kwargs = {"new_x": XPos.LMARGIN, "new_y": YPos.NEXT}
        heading_sizes = {1: (18, 10), 2: (14, 8), 3: (12, 7), 4: (11, 6)}
        for kind, block in _blocks(content):
            if kind == "table":
                self._pdf_table(pdf, _table_rows(block))
                continue
            line = block
            text = self._sanitize_for_pdf(line)
            stripped = text.strip()
            heading = re.match(r"^(#{1,4}) (.*)$", text)
            image = IMAGE_LINE.match(stripped)
            pdf.set_text_color(0, 0, 0)
            if heading:
                size, height = heading_sizes[len(heading.group(1))]
                pdf.set_font("Helvetica", "B", size)
                pdf.multi_cell(0, height, heading.group(2).replace("**", ""), **cell_kwargs)
            elif image:
                path = self._image_path(image.group(2))
                if path.exists():
                    pdf.image(str(path), w=pdf.epw * 0.85, x=pdf.l_margin + pdf.epw * 0.075)
                else:
                    logger.warning("Image %s not found; leaving it out of the pdf", path)
            elif stripped.startswith("> ") or (stripped.startswith("*") and stripped.endswith("*") and not stripped.startswith("**")):
                pdf.set_font("Helvetica", "I", 10)
                pdf.multi_cell(0, 6, stripped.lstrip("> ").strip("*").replace("**", ""), **cell_kwargs)
            elif stripped:
                bullet = text.startswith("- ")
                body = f"  - {text[2:]}" if bullet else text
                pdf.set_font("Helvetica", "", 11)
                if HIGHLIGHT_PATTERN.search(body):
                    pdf.set_text_color(160, 70, 0)  # markers stand out for the user to fill in
                # fpdf2's markdown treats "--" and "__" as styling too, so it is only used where safe.
                markdown = "**" in body and "--" not in body and "__" not in body
                pdf.multi_cell(0, 6, body if markdown else body.replace("**", ""), markdown=markdown, **cell_kwargs)
            else:
                pdf.ln(3)

        path = self.output_dir / filename
        pdf.output(str(path))
        return path

    def _pdf_table(self, pdf: FPDF, rows: List[List[str]]) -> None:
        if not rows:
            return
        widths = TABLE_COL_WIDTHS if rows[0][0] == TABLE_FIRST_HEADER and len(rows[0]) == len(TABLE_COL_WIDTHS) else None
        pdf.set_text_color(0, 0, 0)
        pdf.set_font("Helvetica", "", 7)
        with pdf.table(col_widths=widths, text_align="LEFT", line_height=3.5, first_row_as_headings=True) as table:
            for cells in rows:
                row = table.row()
                for text in cells:
                    row.cell(self._sanitize_for_pdf(text.replace("**", "")))
        pdf.ln(3)

    @staticmethod
    def _sanitize_for_pdf(text: str) -> str:
        # Core Helvetica font only supports latin-1.
        return text.translate(PDF_REPLACEMENTS).encode("latin-1", errors="replace").decode("latin-1")

    def export_bibliography(self, papers: List[Paper], style: str, filename: str) -> Path:
        formatter = CitationFormatter()
        content = formatter.format_bibliography(papers, style=style)
        return self.export_markdown(content, filename)
